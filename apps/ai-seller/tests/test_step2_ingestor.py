"""Шаг 2: приём документов. Дедуп по хешу, предел размера до чтения, форматы."""

import io

import pytest
from sqlalchemy import func, select

from src.db.models import Document, KnowledgeChunk
from src.knowledge import ingestor
from src.knowledge.ingestor import (
    ACCEPTED_FORMATS,
    FileTooLarge,
    UnsupportedFormat,
    check_size,
    extract_text,
    ingest_document,
)

MB = 1024 * 1024

FAQ_MD = """# Правила

Заезд после четырнадцати часов, выезд до двенадцати. Позднее заселение возможно по договорённости.

Стирка: одна загрузка стоит пятьсот тенге, порошок включён в цену.
"""


async def _ingest(session, embedder, source: str, data: bytes, **overrides):
    params = dict(max_bytes=10 * MB, chunk_chars=900, overlap=150, min_chars=80)
    params.update(overrides)
    return await ingest_document(session, embedder, source=source, data=data, **params)


async def _count_chunks(session) -> int:
    return (await session.execute(select(func.count()).select_from(KnowledgeChunk))).scalar_one()


async def test_same_md_twice_is_not_ingested_again(db_session, fake_embedder) -> None:
    data = FAQ_MD.encode("utf-8")
    first = await _ingest(db_session, fake_embedder, "faq.md", data)
    assert first.created is True
    assert first.chunks_added >= 1
    before = await _count_chunks(db_session)

    second = await _ingest(db_session, fake_embedder, "faq-copy.md", data)
    assert second.created is False
    assert second.chunks_added == 0
    assert second.document.id == first.document.id
    assert await _count_chunks(db_session) == before

    docs = (await db_session.execute(select(func.count()).select_from(Document))).scalar_one()
    assert docs == 1


def test_unknown_extension_is_rejected_and_lists_formats() -> None:
    with pytest.raises(UnsupportedFormat) as info:
        extract_text("setup.exe", b"MZ")
    message = str(info.value)
    for ext in (".md", ".txt", ".pdf", ".docx", ".xlsx"):
        assert ext in message
    assert set(ACCEPTED_FORMATS) == {".md", ".txt", ".pdf", ".docx", ".xlsx"}


async def test_unknown_extension_is_rejected_by_ingest(db_session, fake_embedder) -> None:
    with pytest.raises(UnsupportedFormat):
        await _ingest(db_session, fake_embedder, "setup.exe", b"MZ")


def test_check_size_rejects_declared_size_over_limit() -> None:
    with pytest.raises(FileTooLarge):
        check_size(11 * MB, 10 * MB)
    check_size(10 * MB, 10 * MB)  # ровно предел — проходит


async def test_oversized_data_is_rejected_before_parsing(
    db_session, fake_embedder, monkeypatch: pytest.MonkeyPatch
) -> None:
    def _must_not_parse(filename: str, data: bytes) -> str:
        raise AssertionError("extract_text вызван для файла сверх предела")

    monkeypatch.setattr(ingestor, "extract_text", _must_not_parse)
    with pytest.raises(FileTooLarge):
        await _ingest(db_session, fake_embedder, "big.txt", b"x" * 1001, max_bytes=1000)


async def test_txt_in_cp1251_is_decoded(db_session, fake_embedder) -> None:
    text = "Заезд после четырнадцати часов, выезд до двенадцати."
    result = await _ingest(db_session, fake_embedder, "rules.txt", text.encode("cp1251"))
    assert result.created is True
    rows = (await db_session.execute(select(KnowledgeChunk.content))).scalars().all()
    assert any(text in content for content in rows)


def test_extension_case_is_ignored() -> None:
    assert extract_text("RULES.TXT", "Заезд с 14:00.".encode("utf-8")).strip() == "Заезд с 14:00."


async def test_xlsx_sheet_becomes_table_rows(db_session, fake_embedder) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["Тариф", "Цена"])
    sheet.append(["Базовый", 1900])
    sheet.append(["Про", 4900])
    buf = io.BytesIO()
    book.save(buf)

    result = await _ingest(db_session, fake_embedder, "prices.xlsx", buf.getvalue())
    assert result.chunks_added == 2
    chunks = (
        await db_session.execute(
            select(KnowledgeChunk).where(KnowledgeChunk.document_id == result.document.id)
        )
    ).scalars().all()
    assert len(chunks) == 2
    for chunk in chunks:
        assert chunk.chunk_metadata["kind"] == "table_row"
        assert "Тариф" in chunk.content
        assert "Цена" in chunk.content
    contents = " ".join(c.content for c in chunks)
    assert "1900" in contents and "4900" in contents


async def test_docx_paragraphs_are_ingested(db_session, fake_embedder) -> None:
    from docx import Document as DocxDocument

    first = "Заезд после четырнадцати часов, выезд до двенадцати часов дня."
    second = "Стирка: одна загрузка стоит пятьсот тенге, порошок включён в цену."
    doc = DocxDocument()
    doc.add_paragraph(first)
    doc.add_paragraph(second)
    buf = io.BytesIO()
    doc.save(buf)

    result = await _ingest(db_session, fake_embedder, "rules.docx", buf.getvalue())
    assert result.chunks_added >= 1
    contents = " ".join(
        (await db_session.execute(select(KnowledgeChunk.content))).scalars().all()
    )
    assert first in contents
    assert second in contents


def test_pdf_blank_page_gives_empty_text() -> None:
    # Разбор .pdf проверяем на уровне extract_text: pypdf читает то, что сам пишет.
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)
    assert extract_text("blank.pdf", buf.getvalue()).strip() == ""


async def test_empty_txt_creates_document_with_zero_chunks(db_session, fake_embedder) -> None:
    result = await _ingest(db_session, fake_embedder, "empty.txt", b"")
    assert result.created is True
    assert result.chunks_added == 0
    assert result.document.chunk_count == 0
    stored = await db_session.get(Document, result.document.id)
    assert stored is not None
    assert stored.chunk_count == 0


async def test_chunks_are_embedded_in_one_call(db_session, fake_embedder, fake_backend) -> None:
    await _ingest(db_session, fake_embedder, "faq.md", FAQ_MD.encode("utf-8"), chunk_chars=120)
    assert fake_backend.calls == 1
    embedded = (await db_session.execute(select(KnowledgeChunk.embedding))).scalars().all()
    assert embedded
    assert all(vec is not None and len(vec) == 384 for vec in embedded)


def test_pipe_in_cell_survives_markdown_roundtrip() -> None:
    # «|» внутри ячейки docx/xlsx: экранирование ingestor и разбор chunker должны совпадать.
    from src.knowledge.chunker import split_text

    table = ingestor._markdown_table([["A", "B"], ["x|y", "z"]])
    rows = [c for c in split_text(table) if c.meta["kind"] == "table_row"]
    assert len(rows) == 1
    assert rows[0].text == "A: x|y; B: z"
