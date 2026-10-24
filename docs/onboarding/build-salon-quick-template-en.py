from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import sys

# Salon quick-start questionnaire (English, country-neutral payments). Examples are fictional.
OUT = sys.argv[1]
wb = Workbook()
HEAD = PatternFill('solid', fgColor='1F2A44')
EX = Font(italic=True, color='808080')
thin = Side(style='thin', color='D0D5DD')
B = Border(left=thin, right=thin, top=thin, bottom=thin)
YN = ['Yes', 'No']


def table(ws, title, cols, example, rows, dv=None):
    ws.title = title
    for i, (name, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = B
        ws.column_dimensions[get_column_letter(i)].width = width
        e = ws.cell(row=2, column=i, value=example[i - 1])
        e.font = EX
        e.border = B
        for r in range(3, 3 + rows):
            ws.cell(row=r, column=i).border = B
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = 'A3'
    for col, values in (dv or {}).items():
        v = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
        ws.add_data_validation(v)
        v.add(f'{col}3:{col}{2 + rows}')


def qa_sheet(ws, title, intro, items, dv_rows=None):
    ws.title = title
    ws['A1'] = intro
    ws['A1'].font = Font(bold=True, size=12)
    ws['A1'].alignment = Alignment(wrap_text=True)
    ws.merge_cells('A1:C1')
    ws.row_dimensions[1].height = 36
    for i, h in enumerate(['Question', 'Answer', 'Example'], 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD
    for i, (q, ex) in enumerate(items, 3):
        ws.cell(row=i, column=1, value=q).border = B
        ws.cell(row=i, column=2).border = B
        e = ws.cell(row=i, column=3, value=ex)
        e.font = EX
        e.border = B
    ws.column_dimensions['A'].width = 52
    ws.column_dimensions['B'].width = 36
    ws.column_dimensions['C'].width = 40
    if dv_rows:
        v = DataValidation(type='list', formula1='"Yes,No"', allow_blank=True)
        ws.add_data_validation(v)
        for r in dv_rows:
            v.add(f'B{r}')


qa_sheet(wb.active, 'Salon',
         'WETOP salon quick start. Grey text shows examples. Fill in the "Answer" column.',
         [
             ('Salon name', 'Pearl Beauty Lounge'),
             ('City and full address', 'Doha, West Bay, Tower 5, 2nd floor'),
             ('Map link (Google Maps)', 'https://maps.google.com/...'),
             ('Phone for clients', '+974 0000 0000'),
             ('Email for clients', 'hello@pearl-beauty.example'),
             ('Opening hours by day', 'Sat–Thu 10:00–22:00, Fri 14:00–22:00'),
             ('Currency', 'QAR'),
             ('Languages for clients', 'English, Arabic'),
         ])

qa_sheet(wb.create_sheet(), 'Payments',
         'Which payment methods do you accept? Answer Yes or No.',
         [
             ('Cash', 'Yes'),
             ('Card (POS terminal, incl. Apple Pay / Google Pay)', 'Yes'),
             ('Bank transfer to the company account', 'No'),
             ('Bank transfer to a personal account', 'No'),
             ('Other (payment link, wallet app): which one?', 'payment link'),
             ('Gift certificates', 'No'),
             ('Deposit / prepayment for a booking? How much?', 'No'),
             ('Cash in the till on the start day', '500'),
         ], dv_rows=[3, 4, 5, 6, 8])

qa_sheet(wb.create_sheet(), 'Administrator',
         'Administrator who is also the cashier. Each person signs in with their own email.',
         [
             ('Full name', 'Sara Ahmed'),
             ('Email for sign-in', 'sara@pearl-beauty.example'),
             ('Phone', '+974 0000 0001'),
             ('Can refund money to a client?', 'No'),
             ('Can see the revenue of the whole salon?', 'No'),
             ('Owner / manager: full name and email (if they also sign in)', ''),
         ], dv_rows=[6, 7])

table(wb.create_sheet(), 'Specialists',
      [('Full name', 28), ('Email for sign-in (optional)', 30), ('Phone', 18), ('Services they do', 40), ('Working days and hours', 32)],
      ['Maria Santos', 'maria@pearl-beauty.example', '+974 0000 0002', 'haircut, colouring, blow-dry', 'Sat, Mon, Wed 10:00–19:00'],
      25)

table(wb.create_sheet(), 'Services',
      [('Category', 20), ('Service', 34), ('Duration, min', 14), ('Price', 12)],
      ['Hair', "Women's haircut", 60, 150],
      80)

table(wb.create_sheet(), 'Extras',
      [('Item', 30), ('Free or paid', 16), ('Price', 12)],
      ['Arabic coffee', 'Free', ''],
      15, dv={'B': ['Free', 'Paid']})

wb.save(OUT)
print('saved', OUT)
